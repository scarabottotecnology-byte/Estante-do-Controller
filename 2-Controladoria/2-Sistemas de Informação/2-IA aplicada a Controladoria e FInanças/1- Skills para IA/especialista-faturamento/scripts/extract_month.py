#!/usr/bin/env python3
"""Extrai dados de UM arquivo mensal de faturamento e salva em pickle (cache)."""
import sys
import re
import os
import pickle
from datetime import datetime
import pandas as pd
import openpyxl


def parse_filename(fname):
    m = re.match(r"(\d{2})-(\d{2})_-_Faturamento_(\w+)\.xlsx", fname)
    if not m:
        return None
    mes, ano, canal = m.groups()
    return int(mes), 2000 + int(ano), canal.upper()


def read_dashboard(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb['Dashboard']
    total = {
        'faturamento': ws['A4'].value or 0,
        'pedidos': ws['B4'].value or 0,
        'ticket_medio': ws['C4'].value or 0,
        'itens': ws['D4'].value or 0,
    }
    bu = {}
    for r in range(8, 13):
        nome = ws.cell(r, 1).value
        if nome:
            bu[nome] = {
                'faturamento': ws.cell(r, 2).value or 0,
                'pedidos': ws.cell(r, 4).value or 0,
                'itens': ws.cell(r, 6).value or 0,
            }
    canal = {}
    for r in (17, 18):
        nome = ws.cell(r, 1).value
        if nome:
            canal[nome] = {'faturamento': ws.cell(r, 2).value or 0, 'pedidos': ws.cell(r, 4).value or 0}
    wb.close()
    return total, bu, canal


def read_fat_diario(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb['Fat. Diário']
    rows = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row[0] is None:
            continue
        data = row[0]
        if isinstance(data, str):
            try:
                data = datetime.strptime(data.strip(), "%d/%m/%Y")
            except ValueError:
                continue
        if not isinstance(data, datetime):
            continue
        rows.append({
            'data': data,
            'faturamento': row[2] or 0,
            'qtde': row[3] or 0,
            'receb_30d': row[7] or 0,
            'receb_60d': row[8] or 0,
            'receb_90d': row[9] or 0,
        })
    wb.close()
    return pd.DataFrame(rows)


def read_sheet1_agg(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb['Sheet1']
    from collections import defaultdict
    agg = defaultdict(lambda: [0.0, set()])
    for row in ws.iter_rows(min_row=2, values_only=True):
        pedido, total, empresa, bu = row[0], row[5], row[8], row[10]
        if empresa is None or empresa == 0:
            continue
        key = (empresa, bu)
        agg[key][0] += total if isinstance(total, (int, float)) else 0
        agg[key][1].add(pedido)
    wb.close()
    rows = [{'empresa': k[0], 'bu': k[1], 'faturamento': v[0], 'pedidos': len(v[1])} for k, v in agg.items()]
    return pd.DataFrame(rows)


def read_cancelados(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if 'Cancelados' not in wb.sheetnames:
        wb.close()
        return pd.DataFrame(columns=['bu', 'cancelado', 'qtd'])
    ws = wb['Cancelados']
    rows = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row[0] is None:
            continue
        val = row[3] if isinstance(row[3], (int, float)) else 0
        bu_val = row[9] if len(row) > 9 else None
        rows.append({'bu': bu_val, 'total': val})
    wb.close()
    if not rows:
        return pd.DataFrame(columns=['bu', 'cancelado', 'qtd'])
    df = pd.DataFrame(rows)
    agg = df.groupby('bu').agg(cancelado=('total', 'sum'), qtd=('total', 'count')).reset_index()
    return agg


def read_fat_produtos(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if 'Fat. Produtos' not in wb.sheetnames:
        wb.close()
        return pd.DataFrame(columns=['sku', 'produto', 'faturamento', 'qtd', 'cmv'])
    ws = wb['Fat. Produtos']
    rows = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        sku, produto, fat, qtd, cmv = row[1], row[2], row[3], row[4], row[5]
        if not isinstance(fat, (int, float)):
            continue
        rows.append({
            'sku': sku, 'produto': produto, 'faturamento': fat,
            'qtd': qtd if isinstance(qtd, (int, float)) else 0,
            'cmv': cmv if isinstance(cmv, (int, float)) else 0,
        })
    wb.close()
    return pd.DataFrame(rows)


def main(filepath, output_pickle):
    fname = os.path.basename(filepath)
    parsed = parse_filename(fname)
    if not parsed:
        print(f"Arquivo fora do padrão: {fname}")
        sys.exit(1)
    mes, ano, canal = parsed

    total, bu, canal_dash = read_dashboard(filepath)
    diario = read_fat_diario(filepath)
    diario['canal'] = canal
    empresa = read_sheet1_agg(filepath)
    empresa['mes'] = mes
    empresa['ano'] = ano
    empresa['canal'] = canal
    cancelados = read_cancelados(filepath)
    produtos = read_fat_produtos(filepath)
    produtos['mes'] = mes
    produtos['ano'] = ano
    produtos['canal'] = canal

    data = {
        'mes': mes, 'ano': ano, 'canal': canal,
        'total': total, 'bu': bu, 'canal_dash': canal_dash,
        'diario': diario, 'empresa': empresa,
        'cancelados': cancelados, 'produtos': produtos,
    }
    with open(output_pickle, 'wb') as f:
        pickle.dump(data, f)
    print(f"OK: {fname} -> {output_pickle} | fat={total['faturamento']:.2f} pedidos={total['pedidos']}")


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
